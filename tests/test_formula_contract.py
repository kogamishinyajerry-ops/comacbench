"""Spreadsheet contracts tested through the real CLI and LibreOffice conversion."""
from pathlib import Path
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import os
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]

@unittest.skipUnless(shutil.which('soffice'), 'LibreOffice is required for formula-cell integration tests')
class FormulaContractTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='comac-formula-')
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.tasks = self.root/'tasks'/'fixture'
        self.tasks.mkdir(parents=True)
        for name, value in [('init.xlsx',0), ('golden.xlsx',42)]:
            wb = openpyxl.Workbook()
            wb.active.title = 'RequiredSheet'
            wb.active['A1'] = value
            wb.create_sheet('KeepMe')['A1'] = 'metadata'
            wb.save(self.root/name)
            wb.close()
        source = sorted((ROOT/'tasks/spreadsheetbench.verified_subset').glob('*.yaml'))[0]
        spec = yaml.safe_load(source.read_text())
        spec['id'] = 'fixture_001'
        spec['registry_id'] = 'fixture.formula'
        prompt = self.tasks/'fixture_001.md'
        prompt.write_text('Produce output.xlsx. Keep all sheets. Put 42 in RequiredSheet!A1.')
        spec['input'] = {'prompt_file':str(prompt), 'assets':[{'path':str(self.root/'init.xlsx'),
            'digest':hashlib.sha256((self.root/'init.xlsx').read_bytes()).hexdigest()}]}
        spec['grader'].update(answer_sheet='RequiredSheet',answer_position='A1:A1',
            golden_workbook=str(self.root/'golden.xlsx'),oracle_source=str(self.root/'candidate.py'))
        (self.tasks/'fixture_001.yaml').write_text(yaml.safe_dump(spec))

    def evaluate(self, sheet='RequiredSheet', keep=True, value=42, write=True, code=None):
        lines = ['import openpyxl', 'w=openpyxl.Workbook()',f'w.active.title={sheet!r}',
                 f'w.active["A1"]={value!r}']
        if keep:
            lines.append('w.create_sheet("KeepMe")["A1"]="metadata"')
        if write:
            lines.append('w.save("output.xlsx")')
        lines.append('w.close()')
        (self.root/'candidate.py').write_text(code if code is not None else '\n'.join(lines)+'\n')
        out = self.root/'out'
        r = subprocess.run([sys.executable,'-m','runners.design_artifact','--tasks',str(self.tasks),
            '--out',str(out),'--provider','oracle'],cwd=ROOT,capture_output=True,text=True,timeout=60)
        self.assertEqual(r.returncode,0,r.stderr)
        return json.loads((out/'result_fixture_001.json').read_text())

    def deletion_fixture(self):
        for name, rows in [('init.xlsx', ['Header','remove me','also remove']),
                           ('golden.xlsx', ['Header'])]:
            wb = openpyxl.load_workbook(self.root/name)
            ws = wb['RequiredSheet']
            for i, value in enumerate(rows, 1):
                ws.cell(i, 1, value)
            wb.save(self.root/name)
            wb.close()
        path = self.tasks/'fixture_001.yaml'
        spec = yaml.safe_load(path.read_text())
        spec['grader']['answer_position'] = 'A1:B4'
        spec['input']['assets'][0]['digest'] = hashlib.sha256((self.root/'init.xlsx').read_bytes()).hexdigest()
        path.write_text(yaml.safe_dump(spec))

    def test_unchanged_input_does_not_satisfy_required_deletion(self):
        self.deletion_fixture()
        r = self.evaluate(code='import shutil\nshutil.copyfile("init.xlsx", "output.xlsx")\n')
        self.assertEqual(r['validity_gate'], 1)
        self.assertAlmostEqual(r['subscores']['physics'], 1/3, places=4)
        self.assertLess(r['score'], 1)

    def test_correct_deletion_gets_full_credit(self):
        self.deletion_fixture()
        r = self.evaluate(code='import openpyxl\nw=openpyxl.load_workbook("init.xlsx")\n'
            'w["RequiredSheet"].delete_rows(2,2)\nw.save("output.xlsx")\nw.close()\n')
        self.assertEqual(r['validity_gate'], 1)
        self.assertEqual(r['score'], 1)

    def test_generated_oracle_reproduces_blanks_from_gold(self):
        from runners.gen_tasks_office2 import ssb_oracle_script
        self.deletion_fixture()
        code = ssb_oracle_script('init.xlsx', self.root/'golden.xlsx', 'RequiredSheet', 'A1:B4')
        r = self.evaluate(code=code)
        self.assertEqual(r['score'], 1)

    def test_generated_oracle_accepts_sheet_qualified_region(self):
        from runners.gen_tasks_office2 import ssb_oracle_script
        self.deletion_fixture()
        code = ssb_oracle_script('init.xlsx', self.root/'golden.xlsx', 'RequiredSheet', "'RequiredSheet'!A1:B4")
        self.assertEqual(self.evaluate(code=code)['score'], 1)

    def test_partial_deletion_counts_remaining_data_as_wrong(self):
        self.deletion_fixture()
        r = self.evaluate(code='import openpyxl\nw=openpyxl.load_workbook("init.xlsx")\n'
            'w["RequiredSheet"]["A2"]=None\nw.save("output.xlsx")\nw.close()\n')
        self.assertAlmostEqual(r['subscores']['physics'], 2/3, places=4)
        self.assertLess(r['score'], 1)

    def test_extra_content_in_previously_blank_cell_is_penalized(self):
        self.deletion_fixture()
        r = self.evaluate(code='import openpyxl\nw=openpyxl.load_workbook("init.xlsx")\n'
            'w["RequiredSheet"].delete_rows(2,2)\nw["RequiredSheet"]["B4"]="VBA instructions"\n'
            'w.save("output.xlsx")\nw.close()\n')
        self.assertAlmostEqual(r['subscores']['physics'], 2/3, places=4)
        self.assertEqual(json.loads(r['artifacts']['grade_details'])['match']['n_extra'], 1)

    def test_correct_deletion_does_not_excuse_damaged_header(self):
        self.deletion_fixture()
        r = self.evaluate(code='import openpyxl\nw=openpyxl.load_workbook("init.xlsx")\n'
            'w["RequiredSheet"].delete_rows(2,2)\nw["RequiredSheet"]["A1"]="wrong"\n'
            'w.save("output.xlsx")\nw.close()\n')
        self.assertAlmostEqual(r['subscores']['physics'], 2/3, places=4)

    def test_clear_all_can_validly_produce_empty_answer_region(self):
        self.deletion_fixture()
        wb = openpyxl.load_workbook(self.root/'golden.xlsx')
        wb['RequiredSheet']['A1'] = None
        wb.save(self.root/'golden.xlsx')
        wb.close()
        r = self.evaluate(code='import openpyxl\nw=openpyxl.load_workbook("init.xlsx")\n'
            'w["RequiredSheet"].delete_rows(1,3)\nw.save("output.xlsx")\nw.close()\n')
        self.assertEqual(r['validity_gate'], 1)
        self.assertEqual(r['score'], 1)

    def test_single_wrong_cell_cannot_round_up_to_full_credit(self):
        for name in ['init.xlsx', 'golden.xlsx']:
            wb = openpyxl.load_workbook(self.root/name)
            ws = wb['RequiredSheet']
            for row in range(1, 25001):
                ws.cell(row, 1, 'preserve')
            wb.save(self.root/name)
            wb.close()
        path = self.tasks/'fixture_001.yaml'
        spec = yaml.safe_load(path.read_text())
        spec['grader']['answer_position'] = 'A1:A25000'
        spec['input']['assets'][0]['digest'] = hashlib.sha256((self.root/'init.xlsx').read_bytes()).hexdigest()
        path.write_text(yaml.safe_dump(spec))
        r = self.evaluate(code='import openpyxl\nw=openpyxl.load_workbook("init.xlsx")\n'
            'w["RequiredSheet"]["A25000"]="wrong"\nw.save("output.xlsx")\nw.close()\n')
        self.assertLess(r['subscores']['physics'], 1)
        self.assertLess(r['score'], 1)

    def test_wrong_answer_sheet_is_a_contract_failure_even_with_correct_values(self):
        r = self.evaluate(sheet='WrongSheet')
        self.assertEqual(r['validity_gate'],0)
        self.assertEqual(r['score'],0)
        self.assertEqual(r['failure_mode'],'missing_output')

    def test_deleting_an_input_sheet_is_a_contract_failure(self):
        r = self.evaluate(keep=False)
        self.assertEqual(r['validity_gate'],0)
        self.assertEqual(r['score'],0)

    def test_correct_workbook_still_gets_full_credit(self):
        r = self.evaluate()
        self.assertEqual(r['validity_gate'],1)
        self.assertEqual(r['score'],1)

    def test_empty_answer_region_is_missing_output(self):
        r = self.evaluate(value=None)
        self.assertEqual(r['validity_gate'],0)
        self.assertEqual(r['failure_mode'],'missing_output')

    def test_missing_output_file_is_rejected(self):
        r = self.evaluate(write=False)
        self.assertEqual(r['validity_gate'],0)
        self.assertEqual(r['failure_mode'],'missing_output')

    def test_bad_gold_is_a_harness_fault_not_a_model_failure(self):
        wb = openpyxl.Workbook()
        wb.active.title = 'InvalidGold'
        wb.active['A1'] = 42
        wb.save(self.root/'golden.xlsx')
        wb.close()
        r = self.evaluate()
        self.assertEqual(r['failure_mode'],'crash')
        self.assertEqual(r['score'],0)

    def test_glm_model_receives_staged_input_filename_through_real_cli(self):
        requests = []
        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass
            def do_POST(self):
                request = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                requests.append(request)
                prompt = request['messages'][-1]['content']
                filename = 'init.xlsx' if '`init.xlsx`' in prompt else 'input.xlsx'
                code = ('import openpyxl\n'
                        f'w=openpyxl.load_workbook({filename!r})\n'
                        'w["RequiredSheet"]["A1"]=42\n'
                        'w.save("output.xlsx")\nw.close()')
                body = json.dumps({'model':'glm-5.3-flash',
                    'choices':[{'message':{'content':'```python\n'+code+'\n```'}}]}).encode()
                self.send_response(200)
                self.send_header('Content-Length', str(len(body)))
                self.end_headers()
                self.wfile.write(body)
        server = ThreadingHTTPServer(('127.0.0.1',0),Handler)
        threading.Thread(target=server.serve_forever,daemon=True).start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        env = {**os.environ, 'GLM_API_KEY':'local-test-only',
               'GLM_BASE_URL':f'http://127.0.0.1:{server.server_port}/v4'}
        out = self.root/'out'
        command = [sys.executable,'-m','runners.design_artifact','--tasks',str(self.tasks),
                   '--out',str(out),'--provider','glm','--model','glm-5.3-flash','--resume']
        r = subprocess.run(command,cwd=ROOT,env=env,capture_output=True,text=True,timeout=60)
        self.assertEqual(r.returncode,0,r.stderr)
        result = out/'result_fixture_001.json'
        self.assertEqual(json.loads(result.read_text())['score'],1)
        self.assertEqual(requests[0]['model'],'glm-5.3-flash')
        prompt = requests[0]['messages'][-1]['content']
        self.assertIn('`init.xlsx`',prompt)
        self.assertNotIn('golden.xlsx',prompt)
        self.assertNotIn(str(self.root),prompt)
        self.assertIn('execute',prompt.lower())
        self.assertIn('VBA',prompt)
        self.assertIn('unrelated',prompt.lower())
        before = result.read_bytes()
        r = subprocess.run(command,cwd=ROOT,env=env,capture_output=True,text=True,timeout=60)
        self.assertEqual(r.returncode,0,r.stderr)
        self.assertEqual(result.read_bytes(),before)
        self.assertEqual(len(requests),1)
        manifest = json.loads((out/'run_manifest.json').read_text())
        self.assertEqual(manifest['extra']['model'],'glm-5.3-flash')
        self.assertEqual(manifest['extra']['resumed_tasks'],1)

@unittest.skipUnless(shutil.which('soffice'), 'LibreOffice is required')
class FormulaConversionTests(unittest.TestCase):
    def test_official_oracle_with_libreoffice_sheet_rename_keeps_full_credit(self):
        with tempfile.TemporaryDirectory(prefix='comac-formula-conversion-') as temp:
            root = Path(temp)
            tasks = root/'tasks'/'spreadsheetbench.verified_subset'
            tasks.mkdir(parents=True)
            (root/'data').symlink_to(ROOT/'data', target_is_directory=True)
            source = ROOT/'tasks/spreadsheetbench.verified_subset/ssb_22_47.yaml'
            shutil.copy2(source, tasks/source.name)
            shutil.copy2(source.with_suffix('.md'), tasks/source.with_suffix('.md').name)
            out = root/'output'
            r = subprocess.run([sys.executable,'-m','runners.design_artifact','--tasks',str(tasks),
                '--out',str(out),'--provider','oracle'],cwd=ROOT,capture_output=True,text=True,timeout=60)
            self.assertEqual(r.returncode,0,r.stderr)
            row = json.loads((out/'result_ssb_22_47.json').read_text())
            self.assertEqual(row['score'],1,row['artifacts'].get('grade_details'))


if __name__ == '__main__':
    unittest.main()
