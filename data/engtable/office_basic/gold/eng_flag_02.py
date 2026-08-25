"""gold — write the required normalized workbook (STATIC values, no formulas)."""
import csv
import openpyxl

with open("eng_flag_02.csv", newline="", encoding="utf-8") as f:
    rows = list(csv.reader(f))

wb = openpyxl.Workbook()
wb.remove(wb.active)
ws = wb.create_sheet("DATA")
ws.append(['time', 'alt_ft', 'ias_kt', 'load_g', 'status_alt', 'status_ias', 'status_load'])
ws.append(['T000', '10000.0', '140.0', '1.1', 'OK', 'OK', 'OK'])
ws.append(['T001', '10920.0', '149.0', '1.45', 'OK', 'OK', 'OK'])
ws.append(['T002', '11840.0', '158.0', '1.8', 'OK', 'OK', 'OK'])
ws.append(['T003', '14260.0', '167.0', '2.15', 'OK', 'OK', 'OK'])
ws.append(['T004', '13680.0', '176.0', '1.1', 'OK', 'OK', 'OK'])
ws.append(['T005', '14000.0', '185.0', '', 'OK', 'OK', 'MISSING'])
ws.append(['T006', '14920.0', '194.0', '1.8', 'OK', 'OK', 'OK'])
ws.append(['T007', '15840.0', '140.0', '2.15', 'OK', 'OK', 'OK'])
ws.append(['T008', '16760.0', '149.0', '1.1', 'OK', 'OK', 'OK'])
ws.append(['T009', '17680.0', '158.0', '1.45', 'OK', 'OK', 'OK'])
ws.append(['T010', '18000.0', '167.0', '1.8', 'OK', 'OK', 'OK'])
ws.append(['T011', '18920.0', '176.0', '2.15', 'OVER', 'OK', 'OK'])
ws.append(['T012', '19840.0', '185.0', '1.1', 'OVER', 'OK', 'OK'])
ws.append(['T013', '20760.0', '194.0', '1.45', 'OVER', 'OK', 'OK'])
ws.append(['T014', '21680.0', '140.0', '1.8', 'OVER', 'OK', 'OK'])
ws.append(['T015', '22000.0', '204.0', '2.15', 'OVER', 'OK', 'OK'])
ws.append(['T016', '22920.0', '158.0', '1.1', 'OVER', 'OK', 'OK'])
ws = wb.create_sheet("STATS")
ws.append(['channel', 'min', 'max', 'mean'])
ws.append(['alt_ft', 10000.0, 22920.0, 16707.059])
ws.append(['ias_kt', 140.0, 204.0, 167.059])
ws.append(['load_g', 1.1, 2.15, 1.603])
wb.save("output.xlsx")
