"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (1, 1, "DATA 3"),
    (1, 2, 1),
    (2, 1, "DATA 6"),
    (2, 2, 2),
    (3, 1, "DATA 9"),
    (3, 2, 3),
    (4, 1, "DATA 10"),
    (4, 2, 4),
    (5, 1, "DATA 12"),
    (5, 2, 5),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_280_17_init.xlsx")
ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
for cells in ws['A1:B12']:
    for cell in cells:
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = None
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
