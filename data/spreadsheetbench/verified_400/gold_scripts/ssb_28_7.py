"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (1, 1, "400-12345-W"),
    (1, 2, "A-123"),
    (1, 3, "A-235"),
    (2, 2, "A-235"),
    (2, 3, "A-235"),
    (3, 3, "A-235"),
    (4, 3, "A-235"),
    (5, 3, "A-235"),
    (6, 1, "400-12345-D"),
    (6, 2, "X-5556"),
    (6, 3, "X-5556"),
    (7, 3, "X-5556"),
    (8, 3, "X-5556"),
    (9, 3, "X-5556"),
    (10, 3, "X-5556"),
    (11, 3, "X-5556"),
    (12, 1, "400-12345-F"),
    (12, 2, "N-897"),
    (12, 3, "N-897"),
    (13, 3, "N-897"),
    (14, 3, "N-897"),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_28_7_init.xlsx")
ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
