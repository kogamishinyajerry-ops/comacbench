"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (2, 8, 3710),
    (3, 8, 1660),
    (4, 8, 2753),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_263_1_init.xlsx")
ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
for cells in ws['H2:H4']:
    for cell in cells:
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = None
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
