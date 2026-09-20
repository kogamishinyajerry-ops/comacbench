"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (1, 1, "Date"),
    (1, 2, "Branch"),
    (1, 3, "Reference"),
    (1, 4, "Amount"),
    (1, 5, "Match Ref and Val"),
    (2, 1, {"__date__": "2021-06-30 00:00:00"}),
    (2, 2, "BR1 PL"),
    (2, 3, "BR1 Sales JUN"),
    (2, 4, -74.95999999999415),
    (2, 5, 0),
    (3, 3, 31422220),
    (3, 4, 85),
    (3, 5, 1),
    (4, 3, "Total Zero"),
    (4, 4, -74.95999999999415),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_408_5_init.xlsx")
ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
for cells in ws['A1:E137']:
    for cell in cells:
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = None
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
