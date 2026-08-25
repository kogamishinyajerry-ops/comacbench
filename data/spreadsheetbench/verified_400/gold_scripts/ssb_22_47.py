"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (2, 6, 1),
    (2, 7, "HASSONA"),
    (2, 8, 123344555),
    (3, 6, 2),
    (3, 7, "HAMMED"),
    (3, 8, 123344577),
    (4, 6, 3),
    (4, 7, "HASSAN"),
    (4, 8, 123444441),
    (5, 6, 4),
    (5, 7, "HUSSNI"),
    (5, 8, 123688888),
    (6, 6, 5),
    (6, 7, "HASSAN"),
    (6, 8, 133444422),
    (7, 6, 6),
    (7, 7, "HAMUDDA"),
    (7, 8, 133444423),
    (8, 6, 7),
    (8, 7, "HAMAN"),
    (8, 8, 133444424),
    (9, 6, 8),
    (9, 7, "HANA"),
    (9, 8, 133444425),
    (10, 6, 9),
    (10, 7, "HASSNA"),
    (10, 8, 133444441),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_22_47_init.xlsx")
ws = wb["sheet1"] if "sheet1" in wb.sheetnames else wb.worksheets[0]
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
