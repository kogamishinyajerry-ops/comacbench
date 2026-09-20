"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (2, 11, 1),
    (2, 12, "AAADLMN"),
    (2, 14, 2),
    (2, 15, "AAADLMN"),
    (2, 17, 3),
    (2, 18, "AACDELN"),
    (2, 20, 4),
    (2, 21, "AACDELN"),
    (3, 11, 5),
    (3, 12, "AACDELN"),
    (3, 14, 6),
    (3, 15, "AADELMN"),
    (3, 17, 7),
    (3, 18, "AADELNR"),
    (3, 20, 8),
    (3, 21, "AADELNT"),
    (4, 11, 9),
    (4, 12, "AADELNW"),
    (4, 14, 10),
    (4, 15, "AADELNX"),
    (4, 17, 11),
    (4, 18, "AABDILN"),
    (4, 20, 12),
    (4, 21, "AADILNP"),
    (5, 11, 13),
    (5, 12, "AADILNR"),
    (5, 14, 14),
    (5, 15, "AADILNR"),
    (5, 17, 15),
    (5, 18, "AACDLNO"),
    (5, 20, 16),
    (5, 21, "AACDLNO"),
    (6, 11, 17),
    (6, 12, "AADGLNO"),
    (6, 14, 18),
    (6, 15, "AADLMNO"),
    (6, 17, 19),
    (6, 18, "AADLMNO"),
    (6, 20, 20),
    (6, 21, "AADLNOP"),
    (7, 11, 21),
    (7, 12, "AADLNOV"),
    (7, 14, 22),
    (7, 15, "AADLNOZ"),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_290_1_init.xlsx")
ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
for cells in ws['K2:U10']:
    for cell in cells:
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = None
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
