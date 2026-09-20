"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (2, 3, 12),
    (3, 3, 130),
    (4, 3, 90),
    (5, 3, 78),
    (6, 3, 34),
    (7, 3, 66),
    (8, 3, 45),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_267_21_init.xlsx")
ws = wb["SH1"] if "SH1" in wb.sheetnames else wb.worksheets[0]
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
