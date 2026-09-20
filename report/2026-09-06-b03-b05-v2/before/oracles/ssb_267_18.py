"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (1, 7, "[\"/link/\", \"path2img&track/path2img&track\", \"alt-text\", \"path2img&track\"],"),
    (2, 7, "//[\"/link/\", \"path2img&track/path2img&track2\", \"alt-text?\", \"path2img&track2\"],"),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_267_18_init.xlsx")
ws = wb["Folha1"] if "Folha1" in wb.sheetnames else wb.worksheets[0]
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
