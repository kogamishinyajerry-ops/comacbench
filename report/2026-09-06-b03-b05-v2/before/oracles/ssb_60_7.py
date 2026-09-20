"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (3, 1, "Recon"),
    (3, 2, "Recon A"),
    (3, 3, 123),
    (3, 4, "Monthly"),
    (3, 5, "ABC"),
    (4, 1, "Recon"),
    (4, 2, "Recon B"),
    (4, 3, 456),
    (4, 4, "Monthly"),
    (4, 5, "DEF"),
    (5, 1, "Recon"),
    (5, 2, "Recon C"),
    (5, 3, 789),
    (5, 4, "Monthly"),
    (5, 5, "GHI"),
    (6, 1, "MJE"),
    (6, 2, "MJE 1"),
    (6, 3, 987),
    (6, 4, "Daily"),
    (6, 5, "JKL"),
    (7, 1, "Reporting"),
    (7, 2, "Reporting A"),
    (7, 3, 654),
    (7, 4, "Monthly"),
    (7, 5, "MNO"),
    (8, 1, "Recon"),
    (8, 2, "Recon 5"),
    (8, 3, 545),
    (8, 4, "Monthly"),
    (8, 5, "John"),
    (9, 1, "MJE"),
    (9, 2, "MJE 6"),
    (9, 3, 658),
    (9, 4, "As Needed"),
    (9, 5, "Arron"),
    (10, 1, "Other"),
    (10, 2, "Other 7"),
    (10, 3, 124),
    (10, 4, "Quarterly"),
    (10, 5, "Thomas"),
    (11, 1, "Other"),
    (11, 2, "Other 7"),
    (11, 3, 895),
    (11, 4, "Quarterly"),
    (11, 5, "Peter"),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_60_7_init.xlsx")
ws = wb["Consolidated Tracker"] if "Consolidated Tracker" in wb.sheetnames else wb.worksheets[0]
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
