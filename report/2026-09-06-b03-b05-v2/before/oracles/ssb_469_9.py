"""oracle — copy init workbook, write embedded golden answer-region values."""
import datetime as _dt
import json
import openpyxl

VALUES = [
    (1, 8, "Debits"),
    (1, 9, "Credits"),
    (2, 8, 4543.190000000001),
    (3, 8, 28723.24333333334),
    (4, 8, 21846.39666666667),
    (5, 9, 6.933333333333334),
    (6, 8, 3535.258),
    (7, 9, 18867.39133333333),
    (8, 8, 866.2666666666667),
    (9, 9, 11314.198),
    (10, 9, 54666.12466666666),
]


def _v(x):
    if isinstance(x, dict) and "__date__" in x:
        return _dt.datetime.strptime(x["__date__"], "%Y-%m-%d %H:%M:%S") \
            if len(x["__date__"]) > 10 else _dt.datetime.strptime(x["__date__"], "%Y-%m-%d")
    return x


wb = openpyxl.load_workbook("ssb_469_9_init.xlsx")
ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
for row, col, v in VALUES:
    ws.cell(row=row, column=col).value = _v(v)
wb.save("output.xlsx")
