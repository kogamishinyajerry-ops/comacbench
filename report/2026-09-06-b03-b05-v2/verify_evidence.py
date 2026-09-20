"""Offline evidence validation and value/critical-change analysis; no API calls."""
from pathlib import Path
import ast
import base64
import hashlib
import json
import math
import sys
import openpyxl
import yaml

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
sys.path.insert(0,str(ROOT))
from runners.common import write_json_atomic


def read(path):return json.loads(path.read_text())
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def equivalent(actual,expected,tolerance):
    if actual=='':actual=None
    if expected=='':expected=None
    if actual==expected:return True
    if actual is None or expected is None:return False
    if isinstance(actual,bool) or isinstance(expected,bool):return False
    try:
        a,b=float(actual),float(expected)
    except (ValueError,TypeError):return False
    return math.isfinite(a) and math.isfinite(b) and abs(a-b)<=tolerance*max(abs(b),1e-9)


def workbook_diagnostics(directory,task_id,row,preflight):
    task=yaml.safe_load((HERE/'workspace/tasks/spreadsheetbench.verified_subset'/f'{task_id}.yaml').read_text())
    g=task['grader'];path=directory/'workbooks'/task_id/'run1-recalculated.xlsx'
    answer={'task_id':task_id,'gate':row['validity_gate'],'score':row['score'],
            'changed_total':preflight['changed_cell_count'],'changed_matched':None,
            'changed_accuracy':None,'unchanged_damaged':None,'numeric_rel_tol':g['numeric_rel_tol']}
    if not path.exists():return {**answer,'status':'no_recalculated_workbook'}
    wb=openpyxl.load_workbook(path,data_only=True)
    gold=openpyxl.load_workbook(ROOT/g['golden_workbook'],data_only=True)
    if g['answer_sheet'] not in wb.sheetnames:
        wb.close();gold.close();return {**answer,'status':'missing_expected_sheet'}
    got,want=wb[g['answer_sheet']],gold[g['answer_sheet']]
    changed=set(preflight['changed_cells'])
    hits=0;damaged=[];misses=[]
    for cells in want[g['answer_position']]:
        for cell in cells:
            ok=equivalent(got[cell.coordinate].value,cell.value,g['numeric_rel_tol'])
            if cell.coordinate in changed:
                hits+=ok
                if not ok and len(misses)<12:misses.append(cell.coordinate)
            elif not ok:damaged.append(cell.coordinate)
    answer.update(status='measured',changed_matched=hits,changed_accuracy=hits/len(changed),
        unchanged_damaged=len(damaged),changed_miss_examples=misses,unchanged_damage_examples=damaged[:12],
        output_sha256=sha(path))
    wb.close();gold.close()
    return answer


def verify():
    protocol=read(HERE/'protocol.json')
    for name,digest in {**protocol['frozen_files'],**protocol['prior_json_hashes']}.items():
        assert sha(ROOT/name)==digest,name
    assert protocol['budget'] is None
    ids=[tid for g in protocol['groups'] for tid in g['expected_task_ids']]
    assert len(ids)==len(set(ids))==24 and 'ssb_22_47' not in ids
    previous=read(ROOT/'report/2026-09-06-expanded-24-v2/protocol.json')
    assert set(ids)=={t for g in previous['groups'] for t in g['expected_task_ids']}
    assert 'ssb_22_47' in protocol['excluded_tasks']
    for relative,before in read(HERE/'task-specs-before.json').items():
        old,now=yaml.safe_load(before),yaml.safe_load((ROOT/relative).read_text())
        now['input']['prompt_sha256']=old['input']['prompt_sha256']
        assert old==now,relative
    old=json.loads(base64.b64decode((HERE/'before/answers.b64').read_bytes()))
    now=json.loads(base64.b64decode((ROOT/'data/awcom/compliance/hidden/answers.b64').read_bytes()))
    for a,b in zip(old,now):b['prompt_sha256']=a['prompt_sha256']
    assert old==now
    for relative,hashes in read(HERE/'oracle-repair.json').items():
        path=ROOT/relative;before=HERE/'before/oracles'/path.name
        assert sha(path)==hashes['after'] and sha(before)==hashes['before']
        def values(source):
            return next(ast.dump(n.value) for n in ast.parse(source.read_text()).body
                        if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='VALUES' for t in n.targets))
        assert values(path)==values(before),relative
    oracle=read(HERE/'oracle/summary.json')
    oracle_rows=[r for g in oracle['groups'] for r in g['scores']]
    assert oracle['status']=='completed' and oracle['requests']==0 and oracle['model']=='oracle'
    assert len(oracle_rows)==24 and all(r['score']==1 and r['gate']==1 for r in oracle_rows)
    suite_rows=[read(p) for p in (HERE/'oracle-suite').glob('result_*.json')]
    assert len(suite_rows)==25 and all(r['score']==1 and r['validity_gate']==1 for r in suite_rows)
    for group in protocol['groups']:
        if group['name']=='hidden':continue
        for tid in group['expected_task_ids']:
            path=Path('tasks')/group['suite']/(tid+'.yaml')
            original,copy=yaml.safe_load((ROOT/path).read_text()),yaml.safe_load((HERE/'workspace'/path).read_text())
            if group['name']=='formula':
                expected=1e-12 if tid=='ssb_341_40' else 0
                assert copy['grader']['numeric_rel_tol']==copy['reference']['rel_tol']==expected
                copy['grader']['numeric_rel_tol']=original['grader']['numeric_rel_tol']
                copy['reference']['rel_tol']=original['reference']['rel_tol']
            assert original==copy,tid
    preflight={r['task_id']:r for r in read(HERE/'table-preflight.json')}
    output={'status':'passed','protocol_sha256':sha(HERE/'protocol.json'),
        'postfreeze_tasks_unchanged':True,'source_files_verified':len(protocol['frozen_files']),
        'prior_json_artifacts_unchanged':len(protocol['prior_json_hashes']),
        'offline_oracle_full_score':24,'spreadsheet_oracle_full_score':25,'oracle_api_requests':0,'models':{}}
    scores={}
    for alias,config in protocol['models'].items():
        directory=HERE/alias
        summary,ledger=read(directory/'summary.json'),read(directory/'requests.json')
        assert summary['status']=='completed'
        assert ledger['protocol_sha256']==summary['protocol_sha256']==output['protocol_sha256']
        assert len(summary['cache_verification'])==4
        assert all(v['n_cached']==6 and v['results_byte_identical'] and v['mismatched_seed_rejected'] and
                   v['new_http_requests']==0 for v in summary['cache_verification'])
        requests=ledger['requests']
        assert len(requests)>=24 and all(r['status']=='received' for r in requests)
        assert set(r['task_id'] for r in requests)==set(ids)
        public_files=set()
        for request in requests:
            assert request['response_id'] and request['response_model'].lower()==config['model'].lower()
            assert request['request_model']==config['model']
            if request['hidden']:
                assert request['group']=='hidden' and 'public_response_file' not in request
            else:
                path=directory/request['public_response_file'];public_files.add(path)
                assert sha(path)==request['response_sha256']
                response=read(path)
                assert response['model'].lower()==config['model'].lower() and response['id']==request['response_id']
                assert response.get('usage')==request['usage']
        assert public_files==set((directory/'responses').glob('*.json'))
        assert {k:sum((r.get('usage') or {}).get(k,0) for r in requests)
                for k in ('prompt_tokens','completion_tokens','total_tokens')}==summary['usage']
        rows=[];groups=[];diagnostics=[];errors=[]
        for group in protocol['groups']:
            folder=directory/'runs'/group['name']
            manifest=read(folder/'run_manifest.json');identity=manifest['extra']['resume_identity']
            token=hashlib.sha256(json.dumps(identity,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
            current=[read(p) for p in sorted(folder.glob('result_*.json'))]
            assert len(current)==manifest['n_tasks']==manifest['extra']['resumed_tasks']==6
            assert [r['task_id'] for r in current]==group['expected_task_ids']==manifest['extra']['selected_task_ids']
            assert identity['provider']['model']==manifest['extra']['model']==config['model']
            assert identity['provider']['name']==config['provider']
            previous_manifest=read(ROOT/'report/2026-09-06-expanded-24-v2'/alias/'runs'/group['name']/'run_manifest.json')
            assert identity!=previous_manifest['extra']['resume_identity']
            if alias=='glm':
                other=read(HERE/'minimax/runs'/group['name']/'run_manifest.json')
                assert identity!=other['extra']['resume_identity']
            for row in current:
                assert row['artifacts']['resume_identity']==token
                if row['failure_mode']!='crash':
                    assert json.loads(row['artifacts']['model_meta'])['model'].lower()==config['model'].lower()
                if group['name']=='hidden':
                    assert not {'prompt','answer','raw','code'}&row['artifacts'].keys()
                    fields=json.loads(row['artifacts']['layer_details']).get('fields',[])
                    assert all('ref' not in field and 'got' not in field for field in fields)
                if group['name']=='formula':
                    diagnostics.append(workbook_diagnostics(directory,row['task_id'],row,preflight[row['task_id']]))
                elif row['score']<1:
                    details=json.loads(row['artifacts'].get('layer_details','{}'))
                    failed=[f['field'] for f in details.get('fields',[]) if not f.get('ok')]
                    errors.append({'task_id':row['task_id'],'failed_fields':failed,
                        'classification':'identifier_format' if failed==['clause_no'] else 'field_content_or_contract'})
            groups.append({'group':group['name'],'tasks':6,'gate_passed':sum(r['validity_gate']==1 for r in current),
                'full_score':sum(r['score']==1 for r in current),'mean_score':sum(r['score'] for r in current)/6})
            recorded=next(g for g in summary['groups'] if g['name']==group['name'])
            assert recorded['scores']==[{'task_id':r['task_id'],'gate':r['validity_gate'],
                'score':r['score'],'failure_mode':r['failure_mode']} for r in current]
            rows.extend(current)
        assert len(rows)==24
        scores[alias]={r['task_id']:r['score'] for r in rows}
        output['models'][alias]={'model':config['model'],'tasks':24,'requests':len(requests),
            'gate_passed':sum(r['validity_gate']==1 for r in rows),'full_score':sum(r['score']==1 for r in rows),
            'usage':summary['usage'],'groups':groups,'field_diagnostics':errors,'table_diagnostics':diagnostics,
            'cached_results_byte_identical':24,'mismatched_seed_rejected_groups':4,'verification_http_requests':0,
            'public_response_files':len(public_files),'elapsed_s':summary['elapsed_s']}
    offline=read(HERE/'offline/summary.json')
    assert offline['status']=='passed' and offline['api_requests']==0
    assert len(offline['controls'])==6 and len(offline['prior_code_replays'])==12
    assert all(r['correct']==1 and all(r[a]<1 for a in ['unchanged','partial','damaged']) for r in offline['controls'])
    output['offline']={'controls':24,'prior_code_replays':12,'api_requests':0}
    write_json_atomic(HERE/'verification.json',output)
    write_json_atomic(HERE/'scores.json',scores)
    return output,scores


if __name__=='__main__':
    data,scores=verify()
    print(json.dumps({k:{n:m[n] for n in ['tasks','gate_passed','full_score','requests','usage']}
                      for k,m in data['models'].items()},ensure_ascii=False,indent=2))
