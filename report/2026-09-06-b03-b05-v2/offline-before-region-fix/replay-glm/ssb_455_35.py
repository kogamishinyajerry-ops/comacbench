import openpyxl

INPUT_FILE = "ssb_455_35_init.xlsx"
OUTPUT_FILE = "output.xlsx"

# Load the provided input workbook (formulas kept as-is, all sheets preserved)
wb = openpyxl.load_workbook(INPUT_FILE)

# Locate the target worksheet
if "Worksheet1" in wb.sheetnames:
    ws = wb["Worksheet1"]
else:
    ws = wb[wb.sheetnames[0]]

max_row = ws.max_row
max_col = ws.max_column


def starts_with_H_or_A(value):
    """Return True if the first character of the column B value is 'H' or 'A'."""
    if value is None:
        return False
    s = str(value)
    return len(s) > 0 and s[0] in ("H", "A")


# Collect every row whose column B value starts with 'H' or 'A'
rows_to_keep = []
for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
    b_value = row[1] if max_col >= 2 else None  # column B is the 2nd column
    if starts_with_H_or_A(b_value):
        rows_to_keep.append(row)

# Write the kept rows back starting at row 1 (equivalent to deleting the other rows)
for r_idx, row_values in enumerate(rows_to_keep, start=1):
    for c_idx, val in enumerate(row_values, start=1):
        ws.cell(row=r_idx, column=c_idx, value=val)

# Remove the leftover rows at the bottom that are no longer needed
new_last_row = len(rows_to_keep)
if max_row > new_last_row:
    ws.delete_rows(new_last_row + 1, max_row - new_last_row)

# Save the completed workbook as output.xlsx (all other sheets remain untouched)
wb.save(OUTPUT_FILE)