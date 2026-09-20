import openpyxl
import re
from datetime import datetime

wb = openpyxl.load_workbook('ssb_290_27_init.xlsx')
ws = wb['Sheet1']

def is_date_cell(cell):
    """Check if cell contains a date."""
    if isinstance(cell.value, datetime):
        return True
    if cell.number_format:
        fmt = cell.number_format.lower()
        if any(token in fmt for token in ['yy', 'mm', 'dd']):
            if isinstance(cell.value, (int, float)):
                return True
    return False

# Process column B: remove alphabetic characters and spaces from strings, leave dates/numbers alone
for row in range(1, ws.max_row + 1):
    cell = ws.cell(row=row, column=2)
    if cell.value is not None and not is_date_cell(cell):
        if isinstance(cell.value, str):
            new_value = re.sub(r'[A-Za-z\s]', '', cell.value)
            cell.value = new_value

# Write the modified column B values into the answer region A14:I137 (column A of that range)
for r in range(14, 138):
    src = ws.cell(row=r, column=2)
    dst = ws.cell(row=r, column=1)
    dst.value = src.value

wb.save('output.xlsx')