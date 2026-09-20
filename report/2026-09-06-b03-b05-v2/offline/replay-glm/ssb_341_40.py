import openpyxl

# ----------------------------------------------------------------------------
# The task (originally a VBA question): whenever
#   Column A = 'Government'
#   Column B = 'Germany'
#   Column C = 'Carretera'
# the word 'Volkswagen' should be written into column Q of that row.
#
# Here we perform the same conditional fill with openpyxl on the provided
# workbook, preserving every existing sheet, and save the result as
# output.xlsx.
# ----------------------------------------------------------------------------

INPUT_FILE = "ssb_341_40_init.xlsx"
OUTPUT_FILE = "output.xlsx"

# Load the workbook (all sheets are kept)
wb = openpyxl.load_workbook(INPUT_FILE)

# Work on Sheet1 (fall back to the active sheet if the name ever differs)
if "Sheet1" in wb.sheetnames:
    ws = wb["Sheet1"]
else:
    ws = wb.active


def norm(value):
    """Normalize a cell value for robust comparison."""
    return str(value).strip() if value is not None else ""


# Data starts in row 2 (row 1 holds headers). The answer region is A1:Q701,
# so we process data rows 2 through 701 at most.
last_row = min(ws.max_row, 701)

matches = 0
for r in range(2, last_row + 1):
    col_a = norm(ws.cell(row=r, column=1).value)   # Column A
    col_b = norm(ws.cell(row=r, column=2).value)   # Column B
    col_c = norm(ws.cell(row=r, column=3).value)   # Column C

    if (
        col_a.lower() == "government"
        and col_b.lower() == "germany"
        and col_c.lower() == "carretera"
    ):
        ws.cell(row=r, column=17).value = "Volkswagen"  # Column Q
        matches += 1

print(f"Rows updated with 'Volkswagen' in column Q: {matches}")

# Save a copy of the workbook (all original sheets preserved) as output.xlsx
wb.save(OUTPUT_FILE)
print(f"Saved workbook as {OUTPUT_FILE}")