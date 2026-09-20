"""
Automatically delete rows (columns A:E) whose word in column A matches any
word found in column I, shifting the remaining rows up — the programmatic
equivalent of selecting A1:E1 and pressing CTRL + MINUS for every duplicate.

Result is written into 'Sheet1' cells A1:E1102 of output.xlsx
(a copy of the provided input workbook, all other sheets preserved).
"""

from openpyxl import load_workbook

INPUT_FILE = "ssb_23_24_init.xlsx"
OUTPUT_FILE = "output.xlsx"
SHEET_NAME = "Sheet1"
ANSWER_ROWS = 1102          # answer region: A1:E1102

# ---------------------------------------------------------------- load
wb = load_workbook(INPUT_FILE)                      # keeps every existing sheet
ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

n_rows = max(ws.max_row, ANSWER_ROWS)               # never miss data past row 1102

# ------------------------------------------------- 1) words in column I
# Collect every word that appears anywhere in column I (a cell may hold
# several words separated by spaces, commas, semicolons or tabs).
stop_words = set()
for r in range(1, ws.max_row + 1):
    val = ws.cell(row=r, column=9).value            # column I = 9
    if val is None:
        continue
    text = str(val).strip()
    if not text:
        continue
    for token in text.replace(",", " ").replace(";", " ").replace("\t", " ").split():
        stop_words.add(token.strip().lower())

# ------------------------------- 2) keep rows whose A-word is NOT in I
kept = []
for r in range(1, n_rows + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 6)]   # A..E
    a_val = vals[0]

    if a_val is None or str(a_val).strip() == "":
        kept.append(vals)                           # empty row: nothing to delete
        continue

    if str(a_val).strip().lower() in stop_words:
        continue                                    # match in column I -> delete row

    kept.append(vals)

# ----------------------- 3) clear A1:E{n_rows} and shift remaining up
for r in range(1, n_rows + 1):
    for c in range(1, 6):
        ws.cell(row=r, column=c).value = None

for i, vals in enumerate(kept):
    r = i + 1
    for c, v in enumerate(vals):
        if v is not None:
            ws.cell(row=r, column=c + 1).value = v

# ---------------------------------------------------------------- save
wb.save(OUTPUT_FILE)
print(f"Deleted {n_rows - len(kept)} matching row(s); "
      f"{len(kept)} row(s) written to {SHEET_NAME}!A1:E{ANSWER_ROWS} in {OUTPUT_FILE}")