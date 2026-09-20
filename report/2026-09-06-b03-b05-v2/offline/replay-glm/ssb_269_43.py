import openpyxl

INPUT_FILE = "ssb_269_43_init.xlsx"
OUTPUT_FILE = "output.xlsx"


def has_value(cell):
    """Return True if the cell contains an actual value (not empty/whitespace)."""
    if cell is None:
        return False
    v = cell.value
    if v is None:
        return False
    if isinstance(v, str) and v.strip() == "":
        return False
    return True


def delete_rows_i_filled_h_blank(ws):
    """
    Equivalent of the VBA macro:
        delete a row if column I has a value but the same row's column H is blank.
    Rows are collected first, then deleted bottom-up so indices stay valid.
    """
    rows_to_delete = []
    max_row = ws.max_row

    for r in range(1, max_row + 1):
        i_cell = ws.cell(row=r, column=9)  # Column I
        h_cell = ws.cell(row=r, column=8)  # Column H
        if has_value(i_cell) and not has_value(h_cell):
            rows_to_delete.append(r)

    # Delete from the bottom up (like stepping backwards in the VBA loop),
    # so shifting caused by earlier deletions doesn't break later indices.
    for r in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(r)

    return rows_to_delete


def main():
    wb = openpyxl.load_workbook(INPUT_FILE)

    # Choose Sheet1 (fallback to active sheet if somehow missing)
    if "Sheet1" in wb.sheetnames:
        ws = wb["Sheet1"]
    else:
        ws = wb.active

    deleted = delete_rows_i_filled_h_blank(ws)
    print(f"Deleted {len(deleted)} row(s): {deleted}")

    # Preserve all other sheets untouched, save result
    wb.save(OUTPUT_FILE)
    print(f"Saved workbook as {OUTPUT_FILE}")


if __name__ == "__main__":
    main()