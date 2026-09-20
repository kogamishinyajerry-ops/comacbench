import openpyxl

INPUT_FILE = "ssb_290_27_init.xlsx"
OUTPUT_FILE = "output.xlsx"
SHEET_NAME = "Sheet1"
START_ROW = 14
END_ROW = 137
ANSWER_COL = 1  # column A

answer_text = """HOW TO INSTALL AND RUN (Excel 2016):
1. Press Alt+F11 to open the VBA editor.
2. Menu: Insert -> Module, then paste one of the macros below.
3. Press Alt+Q to close the editor, then Alt+F8 and run
   'RemoveLetterCodesColumnB'. (Save the workbook as .xlsm to keep the macro.)

==================================================
OPTION 1 (recommended) - regex based, very robust
==================================================
Sub RemoveLetterCodesColumnB()

    Dim ws As Worksheet
    Dim rng As Range
    Dim cell As Range
    Dim lastRow As Long
    Dim regEx As Object

    Set ws = ActiveSheet
    Set regEx = CreateObject("VBScript.RegExp")

    regEx.Global = True       ' replace every match, not just the first one
    regEx.Pattern = "[A-Z]"   ' match any uppercase letter

    ' Work over the used cells of column B only
    lastRow = ws.Cells(ws.Rows.Count, "B").End(xlUp).Row
    Set rng = ws.Range("B1:B" & lastRow)

    For Each cell In rng
        ' Only text cells are processed. Cells holding real dates
        ' (VarType = vbDate) or true numbers are skipped untouched,
        ' so the date value in the column is never modified.
        If VarType(cell.Value) = vbString And Len(cell.Value) > 0 Then
            ' PID1 -> 1, PID2 -> 2, GG 1 -> 1, ABC123 -> 123
            cell.Value = Trim(regEx.Replace(cell.Value, ""))
        End If
    Next cell

    MsgBox "Finished: letter codes removed from column B.", vbInformation

End Sub

==================================================
OPTION 2 - plain character loop, no regex needed
==================================================
Sub RemoveLetterCodesColumnB_NoRegex()

    Dim rng As Range, cell As Range
    Dim lastRow As Long, i As Long
    Dim s As String, result As String

    lastRow = Cells(Rows.Count, "B").End(xlUp).Row
    Set rng = Range("B1:B" & lastRow)

    For Each cell In rng
        ' vbString = text (gets cleaned); dates/numbers are left alone
        If VarType(cell.Value) = vbString And Len(cell.Value) > 0 Then
            s = cell.Value
            result = ""
            For i = 1 To Len(s)
                If Mid$(s, i, 1) Like "[A-Z]" Then
                    ' skip the uppercase letter
                Else
                    result = result & Mid$(s, i, 1)
                End If
            Next i
            cell.Value = Trim(result)
        End If
    Next cell

End Sub

NOTES:
- VarType(cell.Value) = vbString guarantees that the date value in
  column B (VarType = vbDate) and any pure numbers are never changed.
- The Trim() call removes the leftover space, so "GG 1" becomes "1".
- Works for any 2- or 3-letter uppercase code with or without a
  space before the number, regardless of string length:
  PID1 -> 1, PID2 -> 2, GG 1 -> 1, ABC123 -> 123."""

lines = answer_text.split("\n")
while lines and lines[-1] == "":
    lines.pop()

wb = openpyxl.load_workbook(INPUT_FILE)
ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

written = 0
for i, line in enumerate(lines):
    row = START_ROW + i
    if row > END_ROW:
        break
    ws.cell(row=row, column=ANSWER_COL, value=line)
    written += 1

wb.save(OUTPUT_FILE)
print(f"Wrote {written} answer lines to {SHEET_NAME}!A{START_ROW}:A{START_ROW + written - 1} -> {OUTPUT_FILE}")