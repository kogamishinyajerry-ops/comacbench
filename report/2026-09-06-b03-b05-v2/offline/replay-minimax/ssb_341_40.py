import openpyxl

# Load the input workbook
wb = openpyxl.load_workbook('ssb_341_40_init.xlsx')

# Get Sheet1
ws = wb['Sheet1']

# Determine the data range (A1:Q701 as specified)
# We'll iterate through rows 1 to 701
max_row = 701

# Iterate through rows and check conditions
for row in range(1, max_row + 1):
    a_val = ws.cell(row=row, column=1).value  # Column A
    b_val = ws.cell(row=row, column=2).value  # Column B
    c_val = ws.cell(row=row, column=3).value  # Column C
    
    # Check the three conditions
    if a_val == 'Government' and b_val == 'Germany' and c_val == 'Carretera':
        # Insert 'Volkswagen' into column Q (column 17)
        ws.cell(row=row, column=17).value = 'Volkswagen'

# Save the workbook as output.xlsx
wb.save('output.xlsx')