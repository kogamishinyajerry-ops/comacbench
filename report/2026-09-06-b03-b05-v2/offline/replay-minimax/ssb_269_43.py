from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries

base_dir = Path.cwd()
source_path = base_dir / "ssb_269_43_init.xlsx"
output_path = base_dir / "output.xlsx"

if not source_path.exists():
    raise FileNotFoundError(f"Required workbook not found: {source_path}")

# Load the exact input file while retaining its sheets and links.
workbook = load_workbook(source_path, keep_links=True)

if "Sheet1" not in workbook.sheetnames:
    raise KeyError("The required worksheet 'Sheet1' does not exist.")

sheet = workbook["Sheet1"]

# Remove any existing merged ranges that overlap the requested answer area.
for merged_range in list(sheet.merged_cells.ranges):
    min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
    if min_row <= 7 and max_row >= 1 and min_col <= 9 and max_col >= 1:
        sheet.unmerge_cells(str(merged_range))

# Clear the requested range.
for row in sheet.iter_rows(min_row=1, max_row=7, min_col=1, max_col=9):
    for cell in row:
        cell.value = None
        cell.data_type = "n"

# Presentation styling.
thin_side = Side(style="thin", color="9FBAD0")
border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
title_fill = PatternFill("solid", fgColor="1F4E78")
code_fill = PatternFill("solid", fgColor="EAF2F8")
title_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
code_font = Font(name="Consolas", size=10, color="203864")

for row_number in range(1, 8):
    for column_number in range(1, 10):
        cell = sheet.cell(row=row_number, column=column_number)
        cell.border = border
        cell.fill = title_fill if row_number == 1 else code_fill
        cell.font = title_font if row_number == 1 else code_font
        cell.alignment = (
            Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            if row_number == 1
            else Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
                indent=1,
            )
        )

sheet.row_dimensions[1].height = 38
for row_number in range(2, 8):
    sheet.row_dimensions[row_number].height = 25

# Merge each presentation row across A:I and place the answer below.
for row_number in range(1, 8):
    sheet.merge_cells(
        start_row=row_number,
        start_column=1,
        end_row=row_number,
        end_column=9,
    )

sheet["A1"] = (
    "Save as .xlsm first; then press Alt+F11, choose Insert > Module, "
    "paste the code, and run DeleteRows. The bottom-up loop prevents skipped "
    "rows; change r > 1 if row 1 is data:"
)
sheet["A2"] = "Sub DeleteRows()"
sheet["A3"] = (
    'Dim r As Long, hVal As String, iVal As String: '
    'With ThisWorkbook.Worksheets("Sheet1")'
)
sheet["A4"] = (
    'For r = .Cells(.Rows.Count, "I").End(xlUp).Row To 1 Step -1'
)
sheet["A5"] = (
    'hVal = Trim$(.Cells(r, "H").Value2 & ""): '
    'iVal = Trim$(.Cells(r, "I").Value2 & "")'
)
sheet["A6"] = (
    'If r > 1 And hVal = "" And iVal <> "" Then .Rows(r).Delete'
)
sheet["A7"] = "Next r: End With: End Sub"

sheet.sheet_view.showGridLines = False
workbook.active = workbook.sheetnames.index("Sheet1")

# Save the completed copy as output.xlsx, preserving all existing sheets.
workbook.save(output_path)
workbook.close()