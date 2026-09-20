import openpyxl

wb = openpyxl.load_workbook('ssb_23_24_init.xlsx')
ws = wb['Sheet1']

# Collect all words from column I (column index 9)
words_in_I = set()
for row in range(1, ws.max_row + 1):
    cell_value = ws.cell(row=row, column=9).value
    if cell_value is not None:
        words_in_I.add(str(cell_value).strip().lower())

# Identify rows whose column A value appears in column I
rows_to_delete = []
for row in range(1, ws.max_row + 1):
    cell_a = ws.cell(row=row, column=1).value
    if cell_a is not None:
        a_value = str(cell_a).strip().lower()
        if a_value in words_in_I:
            rows_to_delete.append(row)

# Delete matching rows from bottom to top so indices stay valid
# delete_rows shifts cells up by default (like Ctrl + Minus)
for row in reversed(rows_to_delete):
    ws.delete_rows(row)

wb.save('output.xlsx')