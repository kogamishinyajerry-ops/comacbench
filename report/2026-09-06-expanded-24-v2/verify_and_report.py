"""Offline evidence validation and value/critical-change analysis; no API calls."""
from pathlib import Path
import hashlib
import json
import math
import sys
import openpyxl
import yaml

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
sys.path.insert(0,str(ROOT))
from runners.common import write_json_atomic


def read(path):return json.loads(path.read_text())
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def equivalent(actual,expected,tolerance):
    if actual=='':actual=None
    if expected=='':expected=None
    if actual==expected:return True
    if actual is None or expected is None:return False
    if isinstance(actual,bool) or isinstance(expected,bool):return False
    try:
        a,b=float(actual),float(expected)
    except (ValueError,TypeError):return False
    return math.isfinite(a) and math.isfinite(b) and abs(a-b)<=tolerance*max(abs(b),1e-9)


def workbook_diagnostics(directory,task_id,row,preflight):
    task=yaml.safe_load((HERE/'workspace/tasks/spreadsheetbench.verified_subset'/f'{task_id}.yaml').read_text())
    g=task['grader'];path=directory/'workbooks'/task_id/'run1-recalculated.xlsx'
    answer={'task_id':task_id,'gate':row['validity_gate'],'score':row['score'],
            'changed_total':preflight['changed_cell_count'],'changed_matched':None,
            'changed_accuracy':None,'unchanged_damaged':None,'numeric_rel_tol':g['numeric_rel_tol']}
    if not path.exists():return {**answer,'status':'no_recalculated_workbook'}
    wb=openpyxl.load_workbook(path,data_only=True)
    gold=openpyxl.load_workbook(ROOT/g['golden_workbook'],data_only=True)
    if g['answer_sheet'] not in wb.sheetnames:
        wb.close();gold.close();return {**answer,'status':'missing_expected_sheet'}
    got,want=wb[g['answer_sheet']],gold[g['answer_sheet']]
    changed=set(preflight['changed_cells'])
    hits=0;damaged=[];misses=[]
    for cells in want[g['answer_position']]:
        for cell in cells:
            ok=equivalent(got[cell.coordinate].value,cell.value,g['numeric_rel_tol'])
            if cell.coordinate in changed:
                hits+=ok
                if not ok and len(misses)<12:misses.append(cell.coordinate)
            elif not ok:damaged.append(cell.coordinate)
    answer.update(status='measured',changed_matched=hits,changed_accuracy=hits/len(changed),
        unchanged_damaged=len(damaged),changed_miss_examples=misses,unchanged_damage_examples=damaged[:12],
        output_sha256=sha(path))
    wb.close();gold.close()
    return answer


def verify():
    protocol=read(HERE/'protocol.json')
    for name,digest in {**protocol['frozen_files'],**protocol['prior_json_hashes']}.items():
        assert sha(ROOT/name)==digest,name
    assert protocol['budget'] is None
    ids=[tid for g in protocol['groups'] for tid in g['expected_task_ids']]
    assert len(ids)==len(set(ids))==24 and 'ssb_22_47' not in ids
    previous=read(ROOT/'report/2026-09-06-b05-live/protocol.json')
    assert not set(ids)&{t for g in previous['groups'] for t in g['expected_task_ids']}
    assert 'ssb_22_47' in protocol['excluded_tasks']
    oracle=read(HERE/'oracle/summary.json')
    oracle_rows=[r for g in oracle['groups'] for r in g['scores']]
    assert oracle['status']=='completed' and oracle['requests']==0 and oracle['model']=='oracle'
    assert len(oracle_rows)==24 and all(r['score']==1 and r['gate']==1 for r in oracle_rows)
    for group in protocol['groups']:
        if group['name']=='hidden':continue
        for tid in group['expected_task_ids']:
            path=Path('tasks')/group['suite']/(tid+'.yaml')
            original,copy=yaml.safe_load((ROOT/path).read_text()),yaml.safe_load((HERE/'workspace'/path).read_text())
            if group['name']=='formula':
                expected=1e-12 if tid=='ssb_341_40' else 0
                assert copy['grader']['numeric_rel_tol']==copy['reference']['rel_tol']==expected
                copy['grader']['numeric_rel_tol']=original['grader']['numeric_rel_tol']
                copy['reference']['rel_tol']=original['reference']['rel_tol']
            assert original==copy,tid
    preflight={r['task_id']:r for r in read(HERE/'table-preflight.json')}
    output={'status':'passed','protocol_sha256':sha(HERE/'protocol.json'),
        'canonical_tasks_unchanged':True,'source_files_verified':len(protocol['frozen_files']),
        'prior_json_artifacts_unchanged':len(protocol['prior_json_hashes']),
        'offline_oracle_full_score':24,'oracle_api_requests':0,'models':{}}
    scores={}
    for alias,config in protocol['models'].items():
        directory=HERE/alias
        summary,ledger=read(directory/'summary.json'),read(directory/'requests.json')
        assert summary['status']=='completed'
        assert ledger['protocol_sha256']==summary['protocol_sha256']==output['protocol_sha256']
        assert len(summary['cache_verification'])==4
        assert all(v['n_cached']==6 and v['results_byte_identical'] and v['mismatched_seed_rejected'] and
                   v['new_http_requests']==0 for v in summary['cache_verification'])
        requests=ledger['requests']
        assert len(requests)>=24 and all(r['status']=='received' for r in requests)
        assert set(r['task_id'] for r in requests)==set(ids)
        public_files=set()
        for request in requests:
            assert request['response_id'] and request['response_model'].lower()==config['model'].lower()
            assert request['request_model']==config['model']
            if request['hidden']:
                assert request['group']=='hidden' and 'public_response_file' not in request
            else:
                path=directory/request['public_response_file'];public_files.add(path)
                assert sha(path)==request['response_sha256']
                response=read(path)
                assert response['model'].lower()==config['model'].lower() and response['id']==request['response_id']
                assert response.get('usage')==request['usage']
        assert public_files==set((directory/'responses').glob('*.json'))
        assert {k:sum((r.get('usage') or {}).get(k,0) for r in requests)
                for k in ('prompt_tokens','completion_tokens','total_tokens')}==summary['usage']
        rows=[];groups=[];diagnostics=[];errors=[]
        for group in protocol['groups']:
            folder=directory/'runs'/group['name']
            manifest=read(folder/'run_manifest.json');identity=manifest['extra']['resume_identity']
            token=hashlib.sha256(json.dumps(identity,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
            current=[read(p) for p in sorted(folder.glob('result_*.json'))]
            assert len(current)==manifest['n_tasks']==manifest['extra']['resumed_tasks']==6
            assert [r['task_id'] for r in current]==group['expected_task_ids']==manifest['extra']['selected_task_ids']
            assert identity['provider']['model']==manifest['extra']['model']==config['model']
            assert identity['provider']['name']==config['provider']
            if alias=='glm':
                other=read(HERE/'minimax/runs'/group['name']/'run_manifest.json')
                assert identity!=other['extra']['resume_identity']
            for row in current:
                assert row['artifacts']['resume_identity']==token
                if row['failure_mode']!='crash':
                    assert json.loads(row['artifacts']['model_meta'])['model'].lower()==config['model'].lower()
                if group['name']=='hidden':
                    assert not {'prompt','answer','raw','code'}&row['artifacts'].keys()
                    fields=json.loads(row['artifacts']['layer_details']).get('fields',[])
                    assert all('ref' not in field and 'got' not in field for field in fields)
                if group['name']=='formula':
                    diagnostics.append(workbook_diagnostics(directory,row['task_id'],row,preflight[row['task_id']]))
                elif row['score']<1:
                    details=json.loads(row['artifacts'].get('layer_details','{}'))
                    failed=[f['field'] for f in details.get('fields',[]) if not f.get('ok')]
                    errors.append({'task_id':row['task_id'],'failed_fields':failed,
                        'classification':'identifier_format' if failed==['clause_no'] else 'field_content_or_contract'})
            groups.append({'group':group['name'],'tasks':6,'gate_passed':sum(r['validity_gate']==1 for r in current),
                'full_score':sum(r['score']==1 for r in current),'mean_score':sum(r['score'] for r in current)/6})
            recorded=next(g for g in summary['groups'] if g['name']==group['name'])
            assert recorded['scores']==[{'task_id':r['task_id'],'gate':r['validity_gate'],
                'score':r['score'],'failure_mode':r['failure_mode']} for r in current]
            rows.extend(current)
        assert len(rows)==24
        scores[alias]={r['task_id']:r['score'] for r in rows}
        output['models'][alias]={'model':config['model'],'tasks':24,'requests':len(requests),
            'gate_passed':sum(r['validity_gate']==1 for r in rows),'full_score':sum(r['score']==1 for r in rows),
            'usage':summary['usage'],'groups':groups,'field_diagnostics':errors,'table_diagnostics':diagnostics,
            'cached_results_byte_identical':24,'mismatched_seed_rejected_groups':4,'verification_http_requests':0,
            'public_response_files':len(public_files),'elapsed_s':summary['elapsed_s']}
    negative=read(HERE/'negative-control/summary.json')
    assert negative['api_requests']==0 and negative['status']=='reproduced_scoring_defect'
    assert negative['protocol_sha256']==output['protocol_sha256']
    assert len(negative['results'])==6
    for entry in negative['results']:
        row=read(HERE/'negative-control'/f'result_{entry["task_id"]}.json')
        assert row['score']==entry['score'] and row['validity_gate']==entry['gate']==1
        assert json.loads(row['artifacts']['model_meta'])['model']=='unchanged-input'
        assert json.loads(row['artifacts']['grade_details'])==entry['grade_details']
    assert next(r for r in negative['results'] if r['task_id']=='ssb_455_35')['score']==1
    output['negative_control']={'tasks':6,'api_requests':0,'scoring_defect_reproduced':True,
        'scores':{r['task_id']:r['score'] for r in negative['results']},
        'artifact_sha256':{p.name:sha(p) for p in sorted((HERE/'negative-control').glob('*.json'))}}
    write_json_atomic(HERE/'verification.json',output)
    write_json_atomic(HERE/'scores.json',scores)
    return output,scores


def report(data,scores):
    group_rows=[];task_rows=[];critical=[];usage=[]
    for i,group in enumerate(read(HERE/'protocol.json')['groups']):
        values=[]
        for alias in ['minimax','glm']:
            g=data['models'][alias]['groups'][i]
            values.append(f'{g["gate_passed"]}/6 有效，{g["full_score"]}/6 满分')
        group_rows.append('| '+group['name']+' | '+' | '.join(values)+' |')
        for tid in group['expected_task_ids']:
            task_rows.append(f'| {tid} | {scores["minimax"][tid]:g} | {scores["glm"][tid]:g} |')
    for alias in ['minimax','glm']:
        m=data['models'][alias]
        usage.append(f'| {m["model"]} | {m["requests"]} | {m["usage"]["prompt_tokens"]} | {m["usage"]["completion_tokens"]} | {m["usage"]["total_tokens"]} |')
        for d in m['table_diagnostics']:
            hit=f'{d["changed_matched"]}/{d["changed_total"]}' if d['changed_matched'] is not None else '不可测'
            critical.append(f'| {alias} | {d["task_id"]} | {d["score"]:g} | {hit} | {d["unchanged_damaged"]} |')
    text='''# 排序题隔离后的双模型 24 题扩样

本次每模型 24 道新题，与上轮 8 题无交集；两个模型都使用同一批任务。
排序题 ssb_22_47 经独立审计后排除；其原任务、gold 与历史成绩保持不变。
48 次真实请求全部收到响应。MiniMax 24/24 有效、15/24 满分；GLM 23/24 有效、16/24 满分。
这是冻结规则下的原始成绩。扩样发现其他输入契约歧义及删行判分缺口，评分可信关仍未全部通过，不能用于模型排名。

## 排序题审计结论

- helper 优先与全局 H 升序冲突；输入去重有 10 行，声明区域只覆盖 9 行。
- 本地输入/gold 与上游镜像字节一致；题面及区域沿用上游，本地转换不是根因。
- 原 4/27 命中均是对不同 REF 编号应用 0.5% 容差产生的假命中，严格比较为 0/27。
- 两模型 G:H 前 9 行符合 helper 解释，并修改了 gold 保留的 F 序号；需要先澄清排序及列职责。

完整输入、两种排序结果及离线重放证据见 [独立审计](../2026-09-06-ssb22-audit/README.md)。
未选定一种解释改写 gold，不因现有模型输出而追改历史分数。

## 扩样协议与校准

seed=20260906，H0，无 scaffold、无判分反馈、无按得分补跑。每组 6 题；
公开审定与隐藏题均覆盖 acam/rev/trc 三族各 2 题。
表格覆盖名单删行、空值删行、括号提取、字母清理、条件填值、前缀筛选。
这是预先选择的可核验样本，不代表全部 benchmark 或模型总体能力。

MiniMax-M3 与 GLM-5.3-flash 均显式调用原生 provider，temperature=0，max_tokens=32768。
MiniMax 思考模式沿用默认，GLM thinking enabled。没有总费用、总 token 或请求次数预算上限；
单次等待上限 1800 秒，解析失败沿用 provider 原有有限重试。
MiniMax 相对上轮调整了 provider/单次长度设置，因此不作严格前后增益比较。

6 道表格均由独立输入变换验证 gold，且都有实际需改变内容。
该校准证明选定变换与 gold 一致，不证明自然语言只有一种解释，也不证明 grader 能拒绝未完成操作的答案；后续负控揭示了这一边界。
实验任务副本中 5 道使用零数值容差；ssb_341_40 使用 1e-12 消除金融小数重算尾差。
所有 canonical YAML、gold、原生产 grader、权重均未改动。
首次离线校准的记录器接线错误和零容差尾差问题保留在 [v1 校准目录](../2026-09-06-expanded-24/README.md)，
该阶段没有真实模型请求。v2 的 24/24 oracle 满分、4 项记录器测试通过后才执行真实模型。

## 分组结果

| 组 | MiniMax-M3 | GLM-5.3-flash |
| --- | --- | --- |
'''+ '\n'.join(group_rows)+'''

## 失败归因与新发现

### 表格输入契约

- MiniMax 在 ssb_269_43、ssb_455_35，GLM 在 ssb_290_27，把 VBA 代码/使用说明写入了工作表，未执行要求的数据变换。
  269_43 明确要求执行，因此属于交付不符合要求；290_27、455_35 的 VBA 问法与本地追加的 “Write your answer into ... cells”
  也容易诱导这种行为。应明确输出的是执行后的工作簿，代码说明不应写进数据区域。
- MiniMax 的 ssb_290_27 完成了 10/10 个目标值变更，却额外把 B 列复制到 A 列，破坏 95 个原本无需改变的单元格。
  原始输出和重算输出在答题区值一致，因此该损坏来自生成脚本。
- GLM 的 ssb_279_23 保留了括号，例如 `(tig)`，gold 是 `tig`。题面只说删除括号外的内容，
  未明确删除括号本身；10/10 未命中属于边界语义分歧，不能直接归为不会提取文本。
- GLM 的 ssb_455_35 对包括表头在内的所有行应用筛选，最终输出空表；脚本两次均退出 0，
  被 nonempty answer region gate 拦截。gold 保留表头，但题面没有明确这一例外。
  数据行确实全部被删除，不能将 gate=0 解读为模型连接失败或完全没完成筛选。

### QA 字段口径

- MiniMax 的 awx_ce_03/05/06/07/08 均有 clause_no 不保留 `§` 的问题；awc_trc_01 的 clause_ref 同样缺少 `§`。
- family 失败见 MiniMax awx_ce_05、GLM awx_ce_07：模型返回完整分类文本，参考只接受 `loads` / `control`。
  公开题同时要求 copy verbatim，需要把该字段的抽取范围写清楚。
- GLM 两道公开 ACAM 的 clause_id 返回 “文档号 + 节号”，参考只接受文档号，其他方法字段正确；
  两道隐藏 ACAM 也在 clause_id 失败，但隐藏正文不落盘，不能据此断言其具体字符串错误相同。
  这些成绩保留原判，不把标识符/范围分歧直接当成领域知识错误。

### 删行评分的离线负控

真实调用完成后，额外将 6 道表格的输入原样复制为 output.xlsx，交给当前生产 grader；没有调用模型。
ssb_455_35 **什么都不删也得到 1.0 分**。该 gold 仅表头非空，grader 只比较 gold 非空格，
并固定 extra=0，所以未删除的 93,401 个数据单元格不受罚。这是已复现的判分缺口，优先级高于继续扩样。
另外，原样复制在 ssb_290_27 得 0.98915、ssb_341_40 得 0.9997，进一步说明总分受未变内容主导。
见 [负控结果](negative-control/summary.json) 与 [复现脚本](negative_control.py)。
这是事后诊断，未据此重选样本、重算模型原始成绩或补跑模型。

## 表格关键变更

总分包含大量原本正确的单元格，工作簿基本要求又占当前适用权重的 50%，需结合关键变更命中判断。
例如 GLM ssb_290_27 总分 0.9239，但 10 个待改值均未改；其额外写入的说明有 4 格被 Excel 公式机制解释，重算为错误值。
相比之下，GLM ssb_279_23 的 0/10 是括号边界分歧；GLM ssb_455_35 的数据删除完成但表头也被删除。
关键变更由 input 与 gold 差分预先确定；同时检查未变单元格是否被破坏。
仅覆盖单元格值和现有工作簿 gate，未扩大到所有格式/样式要求。

| 模型 | 题目 | 原权重总分 | 关键变更命中 | 未变单元格损坏 |
| --- | --- | ---: | ---: | ---: |
'''+ '\n'.join(critical)+'''

## 每题得分

| 题目 | MiniMax-M3 | GLM-5.3-flash |
| --- | ---: | ---: |
'''+ '\n'.join(task_rows)+'''

## 真实请求记录

| 模型 | 请求次数 | 输入 token | 输出 token | 合计 token |
| --- | ---: | ---: | ---: | ---: |
'''+ '\n'.join(usage)+'''

[MiniMax ledger](minimax/requests.json) / [GLM ledger](glm/requests.json)记录实际响应模型、ID、哈希和用量。
公开题的完整原始响应位于各模型 responses/；隐藏题只保留必要元数据，正文与答案不落盘。
表格原始和重算工作簿位于各模型 workbooks/。这些是客户端实测证据，不是服务器签名或收费账单。

## 复核与边界

[机器核验结果](verification.json)验证源码/数据哈希、原证据未变、两模型不同身份、
各 24 份同身份结果字节复用、各 4 组错误 seed 拒绝，缓存校验新增 HTTP=0。
两模型请求可并行；LibreOffice 使用跨进程锁串行转换，以避免共享单实例冲突。

```sh
.venv/bin/python report/2026-09-06-expanded-24-v2/test_ledger.py -v
.venv/bin/python report/2026-09-06-expanded-24-v2/verify_and_report.py
git diff --check
```

需要复用缓存时，在冻结源码/任务未改变的前提下执行：

```sh
.venv/bin/python report/2026-09-06-expanded-24-v2/run_experiment.py --execute minimax
.venv/bin/python report/2026-09-06-expanded-24-v2/run_experiment.py --execute glm
```

## 后续优先级

1. B03 单独修复删行/清空的比较范围：gold 空值也属于预期，使用 input/gold 差分识别必须删除的内容；
   先加入“原样复制不得满分”的负控，再改 grader，保留旧成绩并使用新评分身份。
2. B05 单独明确“执行数据操作、只提交处理后工作簿、保留表头与无关列”，以及括号、family、clause_id 的字段范围；
   不混进 B03 判分修复。
3. 将 ssb_22_47、ssb_279_23、ssb_455_35 的上述歧义列入修订清单；排序题先确定主排序、输出行数与 F 列职责，
   使用独立题目版本，不按当前模型答案反向修 gold。
4. 新一轮扩样前，所有表格都通过正确答案、未操作答案、局部完成答案、额外损坏答案四类检查。
后续若要模型比较，应再冻结代表性分层抽样、统一思考策略口径并重复抽样；本次不发布模型排名。
所有变更留在审计/实验目录；本轮未提交、推送或修改生产接口。
'''
    (HERE/'README.md').write_text(text)


if __name__=='__main__':
    data,scores=verify();report(data,scores)
    print(json.dumps({k:{n:m[n] for n in ['tasks','gate_passed','full_score','requests','usage']}
                      for k,m in data['models'].items()},ensure_ascii=False,indent=2))
