"""Input-derived transformations, independent of embedded golden oracle scripts."""
from pathlib import Path
import copy
import hashlib
import json
import re
import openpyxl
import yaml

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
TABLE_IDS = ['ssb_23_24','ssb_269_43','ssb_279_23','ssb_290_27','ssb_341_40','ssb_455_35']


def normal(value):
    return None if value=='' else value


def derive(tid, source):
    rows = copy.deepcopy(source)
    if tid=='ssb_23_24':
        denied = {r[8] for r in rows}-{None,''}
        shifted = [r[:5] for r in rows if r[0] not in denied]
        for i,row in enumerate(rows):
            row[:5] = shifted[i] if i<len(shifted) else [None]*5
    elif tid=='ssb_269_43':
        rows = [r for r in rows if not (normal(r[8]) is not None and normal(r[7]) is None)]
    elif tid=='ssb_279_23':
        for row in rows:
            if isinstance(row[1],str) and (match:=re.search(r'\((.*?)\)',row[1])):
                row[1]=match.group(1)
    elif tid=='ssb_290_27':
        for row in rows:
            if isinstance(row[1],str) and (match:=re.fullmatch(r'[A-Z]{2,3}\s*(\d+)',row[1])):
                row[1]=int(match.group(1))
    elif tid=='ssb_341_40':
        for row in rows[1:]:
            if tuple(row[:3])==('Government','Germany','Carretera'):
                row[16]='Volkswagen'
    elif tid=='ssb_455_35':
        rows = rows[:1]+[r for r in rows[1:] if str(r[1] or '').startswith(('H','A'))]
    else:
        raise ValueError(tid)
    return rows


def evaluate(tid):
    path = ROOT/'tasks/spreadsheetbench.verified_subset'/f'{tid}.yaml'
    task = yaml.safe_load(path.read_text())
    g = task['grader']
    initial,golden = ROOT/task['input']['assets'][0]['path'],ROOT/g['golden_workbook']
    w = openpyxl.load_workbook(initial,data_only=True)
    z = openpyxl.load_workbook(golden,data_only=True)
    sheet,gold = w[g['answer_sheet']],z[g['answer_sheet']]
    source = [list(row) for row in sheet.values]
    expected = derive(tid,source)
    changed, mismatches, nonempty = [],[],0
    for row in sheet[g['answer_position']]:
        for cell in row:
            i,j=cell.row-1,cell.column-1
            want=normal(gold.cell(i+1,j+1).value)
            got=normal(expected[i][j]) if i<len(expected) and j<len(expected[i]) else None
            if want is not None:
                nonempty+=1
            if normal(cell.value)!=want:
                changed.append(cell.coordinate)
            if got!=want:
                mismatches.append(cell.coordinate)
    assert not mismatches,(tid,mismatches[:10])
    assert changed, f'{tid}: no-op task'
    result={'task_id':tid,'status':'passed','answer_sheet':g['answer_sheet'],
        'answer_position':g['answer_position'],'gold_nonempty_cells':nonempty,
        'changed_cell_count':len(changed),'changed_cells':changed,
        'independent_transform_gold_mismatches':0,
        'input_sha256':hashlib.sha256(initial.read_bytes()).hexdigest(),
        'gold_sha256':hashlib.sha256(golden.read_bytes()).hexdigest(),
        'source':str(path.relative_to(ROOT)),
        'scope':'Declared region only; copy/filter/text transforms, no approximate arithmetic.'}
    w.close();z.close()
    return result


if __name__=='__main__':
    rows=[evaluate(tid) for tid in TABLE_IDS]
    (HERE/'table-preflight.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps([{k:r[k] for k in ['task_id','changed_cell_count',
         'independent_transform_gold_mismatches']} for r in rows],ensure_ascii=False))
