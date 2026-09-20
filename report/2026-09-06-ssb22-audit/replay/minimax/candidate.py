import openpyxl
from openpyxl import load_workbook

wb = load_workbook('ssb_22_47_init.xlsx')
ws = wb['sheet1']

max_row = ws.max_row

header_keywords = {'name', 'names', 'header', 'id', 'number', 'value', 'no', 'no.', 'index', 'item'}

def looks_like_header(val):
    if val is None:
        return True
    if isinstance(val, str) and val.strip().lower() in header_keywords:
        return True
    return False

# Read source data from A:C (starting from row 2)
raw_data = []
for row in range(2, max_row + 1):
    val_a = ws.cell(row=row, column=1).value
    val_b = ws.cell(row=row, column=2).value
    val_c = ws.cell(row=row, column=3).value
    
    # Skip empty cells
    if val_b is None or val_b == '':
        continue
    # Skip headers
    if looks_like_header(val_b):
        continue
    
    raw_data.append((row, val_a, val_b, val_c))

# Read helper column J
helper_order = []
helper_seen = set()
for row in range(2, max_row + 1):
    val_j = ws.cell(row=row, column=10).value
    if val_j is None or val_j == '':
        continue
    if looks_like_header(val_j):
        continue
    if val_j not in helper_seen:
        helper_order.append(val_j)
        helper_seen.add(val_j)

# Remove duplicates (defined by identical entries in both B and C)
seen = set()
filtered_data = []
for item in raw_data:
    key = (item[2], item[3])
    if key not in seen:
        seen.add(key)
        filtered_data.append(item)

# Sort based on helper column or alphabetically
if helper_order:
    priority = {name: idx for idx, name in enumerate(helper_order)}
    original_order = {item[0]: idx for idx, item in enumerate(filtered_data)}
    
    def sort_key(item):
        name = item[2]
        if name in priority:
            return (0, priority[name], original_order[item[0]])
        else:
            return (1, 0, original_order[item[0]])
    
    sorted_data = sorted(filtered_data, key=sort_key)
else:
    sorted_data = sorted(filtered_data, key=lambda x: str(x[2]))

# Clear F2:H10
for row in range(2, 11):
    for col in range(6, 9):
        ws.cell(row=row, column=col).value = None

# Write output to F:H (limit 9 rows for F2:H10)
for i, item in enumerate(sorted_data[:9]):
    ws.cell(row=2 + i, column=6, value=item[1])  # F = A value
    ws.cell(row=2 + i, column=7, value=item[2])  # G = B value (name)
    ws.cell(row=2 + i, column=8, value=item[3])  # H = C value

wb.save('output.xlsx')