"""Independent, offline audit of the published sorting contract and gold."""
from pathlib import Path
import argparse
import hashlib
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
import openpyxl
import yaml

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DATA = ROOT/'data/spreadsheetbench/verified_400'
UPSTREAM = DATA/'spreadsheetbench_verified_400'
PRIOR = ROOT/'report/2026-09-06-b05-live'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def region(sheet, low=2, high=11):
    return [list(r) for r in sheet.iter_rows(min_row=low,max_row=high,min_col=6,max_col=8,values_only=True)]


def audit():
    task = yaml.safe_load((ROOT/'tasks/spreadsheetbench.verified_subset/ssb_22_47.yaml').read_text())
    record = next(r for r in json.loads((UPSTREAM/'dataset.json').read_text()) if r['id']=='22-47')
    prompt_path = ROOT/task['input']['prompt_file']
    initial = ROOT/task['input']['assets'][0]['path']
    golden = ROOT/task['grader']['golden_workbook']
    upstream_dir = UPSTREAM/record['spreadsheet_path']
    assert initial.read_bytes()==(upstream_dir/'1_22-47_init.xlsx').read_bytes()
    assert golden.read_bytes()==(upstream_dir/'1_22-47_golden.xlsx').read_bytes()
    assert prompt_path.read_text().split('\n\n## Answer region')[0].strip()==record['instruction'].strip()
    assert record['instruction'].strip()==(upstream_dir/'prompt.txt').read_text().strip()
    assert record['answer_position']==task['grader']['answer_position']=='F2:H10'
    w = openpyxl.load_workbook(initial,data_only=True)
    g = openpyxl.load_workbook(golden,data_only=True)
    sheet,gold = w['sheet1'],g['sheet1']
    seen,unique,source = set(),[],[]
    for row in sheet.iter_rows(min_row=2,min_col=1,max_col=3):
        item,name,value = [c.value for c in row]
        if name is None or name=='NAME':
            continue
        source.append({'row':row[0].row,'item':item,'name':name,'value':value})
        if (name,value) not in seen:
            unique.append((name,value))
            seen.add((name,value))
    helpers = [sheet.cell(i,10).value for i in range(2,sheet.max_row+1)
               if sheet.cell(i,10).value]
    helper_first = [p for name in helpers for p in unique if p[0]==name]
    helper_first += [p for p in unique if p[0] not in helpers]
    numeric_first = sorted(unique,key=lambda pair:pair[1])
    actual_gold = region(gold)
    gold_pairs = [tuple(row[1:]) for row in actual_gold]
    assert len(unique)==10
    assert gold_pairs==numeric_first
    assert gold_pairs!=helper_first
    # Minimal counterexample: mutually incompatible first-row requirements.
    assert helper_first[0]==('HASSAN',133444422)
    assert gold_pairs[0]==('HASSONA',123344555)
    assert all(gold.cell(i,6).value==sheet.cell(i,6).value==i-1 for i in range(2,12))
    result = {'task_id':'ssb_22_47','api_requests':0,'status':'quarantined',
        'classification':['contradictory_sort_precedence','answer_region_omits_last_output_row',
                          'numeric_tolerance_accepts_different_identifiers'],
        'upstream_mirror':{'prompt_verbatim':True,'input_bytes_equal':True,'gold_bytes_equal':True,
                           'answer_position_unchanged':True},
        'source_rows':source,'helpers':helpers,'unique_pairs':unique,
        'helper_first_output':helper_first,'global_numeric_output':numeric_first,
        'gold_output_F2_H11':actual_gold,
        'gold_matches_helper_first':False,'gold_matches_global_numeric':True,
        'declared_region':'F2:H10','declared_output_rows':9,'actual_unique_rows':10,
        'unscored_nonempty_output':'G11:H11','gold_preserves_prefilled_F':True,
        'replayed_models':{},'source_sha256':{str(p.relative_to(ROOT)):sha(p) for p in
             [initial,golden,prompt_path,UPSTREAM/'dataset.json',upstream_dir/'prompt.txt',
              ROOT/'tasks/spreadsheetbench.verified_subset/ssb_22_47.yaml']}}
    for alias,folder in [('minimax','minimax'),('glm','glm-recovery')]:
        evidence = PRIOR/folder/'runs/formula/result_ssb_22_47.json'
        stored = json.loads(evidence.read_text())
        code = json.loads(stored['artifacts']['code'])
        target = HERE/'replay'/alias
        target.mkdir(parents=True,exist_ok=True)
        code_path = target/'candidate.py'
        code_path.write_text(code)
        with tempfile.TemporaryDirectory(prefix='ssb22-replay-') as temp:
            work = Path(temp)
            shutil.copy2(initial,work/initial.name)
            proc = subprocess.run([sys.executable,'-I',str(code_path)],cwd=work,
                env={'PATH':os.defpath,'HOME':temp},capture_output=True,text=True,timeout=60)
            assert proc.returncode==0,proc.stderr
            output = target/'output.xlsx'
            shutil.copy2(work/'output.xlsx',output)
        cw = openpyxl.load_workbook(output,data_only=True)
        rows = region(cw['sheet1'])
        per_column = {col:sum(rows[i][j]==actual_gold[i][j] for i in range(9))
                      for j,col in enumerate('FGH')}
        tolerant = {col:sum((rows[i][j]==actual_gold[i][j]) or
            (isinstance(rows[i][j],(int,float)) and isinstance(actual_gold[i][j],(int,float)) and
             math.isclose(rows[i][j],actual_gold[i][j],rel_tol=0.005)) for i in range(9))
            for j,col in enumerate('FGH')}
        assert sum(per_column.values())==0
        assert sum(tolerant.values())==4
        assert [tuple(row[1:]) for row in rows[:9]]==helper_first[:9]
        result['replayed_models'][alias] = {'source_result_sha256':sha(evidence),
            'code_sha256':sha(code_path),'output_sha256':sha(output),'rows_F2_H11':rows,
            'GH_matches_helper_first_within_declared_region':True,'GH_matches_global_numeric':False,
            'F_preserved':all(rows[i][0]==i+1 for i in range(10)),
            'exact_matches_in_declared_region':sum(per_column.values()),'exact_matches_by_column':per_column,
            'tolerant_matches_by_column':tolerant,'numeric_identifier_tolerance_false_matches':4,
            'recorded_original_score':stored['score'],'offline_replay_only':True}
        cw.close()
    w.close();g.close()
    (HERE/'audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:result[k] for k in ['status','upstream_mirror','classification',
                                          'declared_output_rows','actual_unique_rows']},ensure_ascii=False))
    return result


if __name__=='__main__':
    p=argparse.ArgumentParser()
    p.add_argument('--assert-helper-contract',action='store_true')
    args=p.parse_args()
    result=audit()
    if args.assert_helper_contract:
        assert result['gold_matches_helper_first'], 'Gold violates helper-first wording; canonical task remains quarantined.'
